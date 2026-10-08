# -*- coding: utf-8 -*-
"""stats.json → Excel 明细报告

用法: python render_excel.py --stats reports/stats.json --out 输出.xlsx
"""
import sys, os, json, argparse
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C
C.setup_stdout()
C.add_site_packages()
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = '微软雅黑'
HDR_FILL = PatternFill('solid', fgColor='1F4E79')
HDR_FONT = Font(name=FONT, size=10, bold=True, color='FFFFFF')
TITLE = Font(name=FONT, size=14, bold=True, color='C4620A')
SUB = Font(name=FONT, size=10, bold=True, color='1F4E79')
CELL = Font(name=FONT, size=10)
NOTE = Font(name=FONT, size=9, color='595959')
FILL_NEG = PatternFill('solid', fgColor='FFC7CE')
FILL_NEU = PatternFill('solid', fgColor='EDEDED')
FILL_POS = PatternFill('solid', fgColor='C6EFCE')
FILL_ALT = PatternFill('solid', fgColor='F2F7FB')
THIN = Side(style='thin', color='BFBFBF')
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CT = Alignment(horizontal='center', vertical='center')
LT = Alignment(horizontal='left', vertical='center', wrap_text=True)

COLNAME = {'date': '数据日期', 'model': '车型', 'label_type': '标签类型', 'l1': '一级标签',
           'l2': '二级标签', 'l3': '三级标签', 'l4': '四级标签', 'vp': '观点标签',
           'emotion': '情感', 'channel': '渠道', 'store': '专营店', 'region': '大区',
           'vin': 'VIN码', 'ccode': '车型码', 'uid': '声音ID', 'text': '内容', 'source_ref': '来源定位',
           'spill_class': '外溢层级', 'matched_terms': '命中词'}

def table(ws, r0, headers, data, widths=None, wrap=(), fill_col=None):
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CT, BD
    for i, row in enumerate(data):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + 1 + i, j, v)
            c.font, c.border = CELL, BD
            c.alignment = LT if j in wrap else CT
            if i % 2 == 1:
                c.fill = FILL_ALT
            if fill_col and j == fill_col:
                if v == '负面':
                    c.fill = FILL_NEG
                elif v == '中性':
                    c.fill = FILL_NEU
                elif v == '正面':
                    c.fill = FILL_POS
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return r0 + 1 + len(data)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--stats', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    S = json.load(open(a.stats, encoding='utf-8'))
    M, m1, m2, m4, m5, m6, m7, m8 = (S['meta'], S['m1'], S['m2'], S['m4'], S['m5'],
                                     S['m6'], S['m7'], S['m8'])
    m3 = S.get('m3')
    m9 = S.get('m9') or {'insights': [], 'actions': [], 'summary': ''}
    topic = M['topic']
    wb = Workbook()

    # ⓪ 关键洞察与建议动作（放最前——这是决策层最先要看的东西）
    ws0 = wb.active
    ws0.title = '⓪ 关键洞察'
    ws0['A1'] = f'{topic} 关键洞察与建议动作'
    ws0['A1'].font = Font(name=FONT, size=16, bold=True, color='C4620A')
    ws0['A2'] = m9.get('summary', '')
    ws0['A2'].font = Font(name=FONT, size=11, bold=True, color='B33A2B')
    LEVEL_CN = {'red': '高风险', 'amber': '需关注', 'info': '提示', 'good': '亮点'}
    ins_rows = []
    for it in m9.get('insights', []):
        ins_rows.append([
            f"{it['priority']}｜{LEVEL_CN.get(it['level'], '')}",
            it['kind'], it['title'],
            '\n'.join(it.get('evidence') or []),
            it['root'], it['action'], it['owner'],
        ])
    table(ws0, 4, ['优先级', '性质', '结论', '支撑事实', '待验证方向', '建议动作', '牵头'],
          ins_rows, widths=[13, 14, 46, 42, 40, 46, 16], wrap=(3, 4, 5, 6, 7))
    for i, it in enumerate(m9.get('insights', [])):
        cell = ws0.cell(5 + i, 1)
        if it['level'] == 'red':
            cell.font = Font(name=FONT, size=10, bold=True, color='B33A2B')
        elif it['level'] == 'amber':
            cell.font = Font(name=FONT, size=10, bold=True, color='C4620A')
        elif it['level'] == 'good':
            cell.font = Font(name=FONT, size=10, bold=True, color='079781')
    # 动作清单另起一段，方便直接复制派活
    r0 = 6 + len(ins_rows)
    ws0.cell(r0, 1, '建议动作清单（按优先级）').font = SUB
    act_rows = [[a['priority'], a['action'], a['owner'], a['from']] for a in m9.get('actions', [])]
    table(ws0, r0 + 1, ['优先级', '动作', '牵头', '来自洞察'], act_rows,
          widths=[13, 60, 16, 14], wrap=(2,))
    ws0.freeze_panes = 'A5'

    # ① 摘要
    ws = wb.create_sheet('① 分析摘要')
    ws['A1'] = f'{topic} 专项简报摘要'
    ws['A1'].font = Font(name=FONT, size=16, bold=True, color='C4620A')
    ws['A2'] = (f"{M['period_name']}｜口径 {M['scope_desc']['main']}｜车型 {M['scope_desc']['target_model']}"
                f"｜{M['date_from']} ~ {M['date_to']}（{M['days']} 天）｜数据 {M['total_rows']:,} 条")
    ws['A2'].font = NOTE
    pos_hint = '、'.join(str(r[0]) for r in (m2.get('vp_top_pos') or [])[:2] if r and r[0])
    rows_sum = [
        ('一、规模', '', ''),
        ('主口径声量', m2['n'], f"渗透率万分之 {m2['penetration']}（分母 {m2['universe_label']} {m2['universe']:,} 条）"),
        ('标签口径对照', m2['robust_n'], '仅按标签命中，用于口径稳健性对照'),
        ('正面', f"{m2.get('positive_n', 0)} 条（{m2.get('positive_rate', 0)}%）",
         f"口碑面{'，高频好评点：' + pos_hint if pos_hint else ''}"),
        ('中性', f"{m2.get('neutral_n', 0)} 条（{m2.get('neutral_rate', 0)}%）", '咨询/期望类，藏着产品机会'),
        ('负面', f"{m2['negative_n']} 条（{m2['negative_rate']}%）", '风险面'),
    ]
    if m3:
        rows_sum += [
            ('', '', ''),
            ('二、基线对比', '', ''),
            (f"{m3['target_label']} 渗透率", f"万分之 {m3['target_rate']}", f"{m3['target_n']} / {m3['target_universe']:,}"),
            (f"{m3['baseline_label']} 渗透率", f"万分之 {m3['baseline_rate']}", f"{m3['baseline_n']} / {m3['baseline_universe']:,}"),
            ('强度倍数', f"{m3['multiple']} 倍" if m3['multiple'] else '—', f"基数差异 {m3['base_gap']}%"),
        ]
    rows_sum += [('', '', ''), ('三、观点主题归并（三档分开）', '', '')]
    rows_sum.append(('负面风险主题规则覆盖', f"{m4.get('shape_coverage', 0)}%",
                     f"已归并 {m4.get('classified_n', 0)} / 负面 {m4.get('negative_n', 0)} 条"))
    if m4.get('rule_review_required'):
        pending = '、'.join(f'{x[0]}({x[1]})' for x in (m4.get('unclassified_vp_top') or [])[:5]) or '无'
        rows_sum.append(('待补规则高频观点', pending,
                         f"未覆盖 {m4.get('unclassified_n', 0)} 条；不展示负面主题图，也不输出主题集中结论"))
    for k, v, p in m4['exclusive']:
        rows_sum.append((k, f'{v} 条', f'{p}%'))
    for emo in ('正面', '中性'):
        theme = (m4.get('themes') or {}).get(emo) or {}
        rows_sum.append((f'{emo}主题规则覆盖', f"{theme.get('coverage', 0)}%",
                         f"已归并 {theme.get('classified_n', 0)} / {emo} {theme.get('total', 0)} 条；不与负面主题混算"))
    rows_sum += [('', '', ''), ('四、趋势（{unit}级）'.format(unit=m5['unit']), '', '')]
    for r in m5['rows']:
        rows_sum.append((f"{r['label']}（{r['range']}）", f"{r['n']} 条",
                         f"{r['days']}天｜均值 {r['avg']}｜占全量万分之 {r['pct']}"))
    rows_sum += [
        ('', '', ''),
        ('五、售后外溢', '', ''),
        ('外溢直命中', m6['spill_total'], '关键词命中、未归入主口径；不含关联扩展'),
        ('其中服务侧（直命中）', m6['service_total'], f"每 1 条主口径伴随 {m6['ratio']} 条"),
        ('关联扩展（不并入直命中）', m6.get('related_total', 0), '仅用户明确确认关联词后输出；不参与服务外溢洞察'),
    ]
    rows_sum += [('', '', ''), ('六、集中度', '', '')]
    if 'vin' in m7:
        rows_sum.append(('重复报修', f"{m7['vin']['repeat']} 台", f"有效 VIN {m7['vin']['valid']}，占 {m7['vin']['rate']}%"))
    if m7.get('专营店'):
        rows_sum.append(('问题最集中专营店', m7['专营店'][0][0], f"{m7['专营店'][0][1]} 条"))

    r = 4
    for j, h in enumerate(['项目', '数值', '说明'], 1):
        c = ws.cell(r, j, h)
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, CT, BD
    r += 1
    for row in rows_sum:
        if row[0] == '' and row[1] == '':
            r += 1
            continue
        if row[1] == '':
            c = ws.cell(r, 1, row[0])
            c.font = SUB
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            for j in (1, 2, 3):
                ws.cell(r, j).fill = PatternFill('solid', fgColor='DDEBF7')
                ws.cell(r, j).border = BD
        else:
            ws.cell(r, 1, row[0]).font = CELL
            ws.cell(r, 1).alignment = LT
            ws.cell(r, 2, row[1]).font = CELL
            ws.cell(r, 2).alignment = CT
            ws.cell(r, 3, row[2]).font = NOTE
            ws.cell(r, 3).alignment = LT
            for j in (1, 2, 3):
                ws.cell(r, j).border = BD
        r += 1
    if M.get('caveats'):
        r += 1
        ws.cell(r, 1, '注意').font = Font(name=FONT, size=10, bold=True, color='C00000')
        ws.cell(r, 2, '；'.join(M['caveats'])).font = Font(name=FONT, size=10, color='C00000')
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 66

    # ② 主口径明细
    det = S.get('detail') or []
    ws2 = wb.create_sheet(f'② 主口径明细({len(det)}条)')
    ws2['A1'] = f'{topic} 主口径明细'
    ws2['A1'].font = TITLE
    keys = list(det[0].keys()) if det else []
    table(ws2, 3, [COLNAME.get(k, k) for k in keys],
          [[d.get(k, '') for k in keys] for d in det],
          widths=[11, 8, 11, 12, 12, 12, 12, 22, 10, 13, 22, 12, 22, 20, 32, 90],
          wrap=(len(keys),), fill_col=keys.index('emotion') + 1 if 'emotion' in keys else None)
    ws2.freeze_panes = 'A4'

    # ③ 观点标签分布（含三档情感构成——每个观点的正负面一眼可见）
    ws3 = wb.create_sheet('③ 观点标签分布')
    ws3['A1'] = '观点标签分布'
    ws3['A1'].font = TITLE
    vp_emo = m4.get('vp_emo') or {}
    rows3 = []
    for k, v in m4['vp_top']:
        ve = vp_emo.get(k) or {}
        rows3.append([k, v, ve.get('正面', 0), ve.get('中性', 0), ve.get('负面', 0)])
    rh = table(ws3, 3, ['观点标签', '声量', '正面', '中性', '负面'], rows3,
               widths=[44, 10, 9, 9, 9])
    rh += 1
    themes = m4.get('themes') or {
        '负面': {
            'total': m4.get('negative_n', 0), 'classified_n': m4.get('classified_n', 0),
            'coverage': m4.get('shape_coverage', 0), 'exclusive': m4.get('exclusive') or [],
            'rule_review_required': m4.get('rule_review_required', False),
            'unclassified_vp_top': m4.get('unclassified_vp_top') or [],
        }
    }
    for emotion in ('负面', '正面', '中性'):
        theme = themes.get(emotion) or {}
        total = theme.get('total', 0)
        classified = theme.get('classified_n', 0)
        coverage = theme.get('coverage', 0)
        if emotion == '负面':
            label = '负面风险主题'
        elif emotion == '正面':
            label = '正面体验主题'
        else:
            label = '中性咨询主题'
        ws3.cell(rh, 1, f'{label}（规则覆盖 {coverage:g}%：{classified}/{total}）').font = SUB
        if theme.get('rule_review_required'):
            pending = '、'.join(f'{x[0]}({x[1]})' for x in (theme.get('unclassified_vp_top') or [])[:5]) or '无'
            rows_theme = [['待补规则高频观点', pending, '覆盖不足，不作主题结构判断']]
        else:
            rows_theme = [[k, v, f'占{emotion} {p}%'] for k, v, p in (theme.get('exclusive') or [])]
        rh = table(ws3, rh + 1, ['主题', '声量', '说明'], rows_theme, widths=[44, 10, 30]) + 1

    # ④ 周期趋势
    ws4 = wb.create_sheet('④ 周期趋势')
    ws4['A1'] = f'周期趋势（{m5["unit"]}级，均值已归一）'
    ws4['A1'].font = TITLE
    rows4 = [[r['label'], r['range'], r['days'], r['total'], r['n'], r['avg'], f"万分之 {r['pct']}"]
             for r in m5['rows']]
    rows4.append(['合计', f"{M['date_from']} ~ {M['date_to']}", M['days'],
                  sum(r['total'] for r in m5['rows']), sum(r['n'] for r in m5['rows']), '', ''])
    table(ws4, 3, ['周期', '日期区间', '天数', '全量声音', f'主口径声量', '周期均值', '占全量比'],
          rows4, widths=[10, 26, 8, 12, 14, 11, 14])

    # ⑤ 外溢明细
    sp = S.get('spill_detail') or []
    direct_n, related_n = m6.get('spill_total', 0), m6.get('related_total', 0)
    ws5 = wb.create_sheet(f'⑤ 外溢明细(直命中{direct_n}条｜关联{related_n}条)')
    ws5['A1'] = '外溢明细（直命中与关联扩展分层，禁止混算）'
    ws5['A1'].font = TITLE
    contract = m6.get('contract') or {}
    ws5['A2'] = ('直命中词：' + ('／'.join(contract.get('direct_terms') or []) or '未配置（外溢已禁用）')
                 + '；关联词：' + ('／'.join(contract.get('related_terms') or []) or '无'))
    ws5['A2'].font = NOTE
    skeys = list(sp[0].keys()) if sp else []
    table(ws5, 4, [COLNAME.get(k, k) for k in skeys],
          [[d.get(k, '') for k in skeys] for d in sp],
          widths=[16, 22, 11, 8, 12, 14, 16, 20, 24, 10, 22, 100],
          wrap=(len(skeys),), fill_col=skeys.index('emotion') + 1 if 'emotion' in skeys else None)
    ws5.freeze_panes = 'A5'

    # ⑥ 重复报修
    ws6 = wb.create_sheet('⑥ 重复报修')
    ws6['A1'] = '重复报修车辆'
    ws6['A1'].font = TITLE
    if 'vin' in m7:
        ws6['A2'] = f"有效 VIN {m7['vin']['valid']} 台，重复 {m7['vin']['repeat']} 台（{m7['vin']['rate']}%）"
        ws6['A2'].font = NOTE
        table(ws6, 4, ['VIN码', '反馈次数', '涉及观点标签'],
              [[a, b, c] for a, b, c in m7['vin']['detail']], widths=[24, 10, 70], wrap=(3,))
    else:
        ws6['A2'] = '数据缺少 VIN 字段，本模块未执行'
        ws6['A2'].font = NOTE

    # ⑦ 集中度分布
    ws7 = wb.create_sheet('⑦ 集中度分布')
    ws7['A1'] = '集中度分布'
    ws7['A1'].font = TITLE
    r7 = 3
    for k in ('大区', '渠道', '专营店'):
        if m7.get(k):
            ws7.cell(r7, 1, f'按{k}').font = SUB
            r7 = table(ws7, r7 + 1, [k, '声量'], [[a, b] for a, b in m7[k]], widths=[28, 14, 18, 18]) + 1
    if m7.get('missing'):
        ws7.cell(r7 + 1, 1, '字段缺失，未执行：' + '、'.join(m7['missing'])).font = NOTE

    # ⑧ 口径与方法说明
    ws8 = wb.create_sheet('⑧ 口径与方法说明')
    ws8['A1'] = '口径、方法与数据局限'
    ws8['A1'].font = TITLE
    notes = [
        ('分析主题', M['query_raw']),
        ('报告周期', M['period_name']),
        ('主口径', M['scope_desc']['main']),
        ('车型范围', M['scope_desc']['target_model']),
        ('基线车型', M['scope_desc']['baseline'] or '未指定（未做基线对比）'),
        ('纳入的容错词', '、'.join(M['scope_desc']['extra_terms']) or '无'),
        ('数据文件', M['data_file']),
        ('数据区间', f"{M['date_from']} ~ {M['date_to']}（{M['days']} 天，共 {M['total_rows']:,} 条）"),
        ('数据体检', f"{m1['dup_desc']}"),
        ('字段缺失率', '；'.join(f'{k} {v}' for k, v in m1['missing'].items())),
        ('生成时间', M['generated_at']),
        ('', ''),
        ('方法局限', ''),
        ('声量≠故障率', 'VOC 声量受客户活跃度、渠道触达、录入习惯影响，不等同于实车故障率。本报告以基线车型做归一化，部分抵消该偏差。'),
        ('归因层级', '统计结论为声量层面的相关性证据，不构成因果结论。根因需工程侧拆解验证。'),
        ('数据边界', '缺少 DMS 工单、返件检测与生产批次数据，无法确认失效机理与受影响批次范围。'),
    ]
    if M.get('caveats'):
        notes.append(('本次注意事项', '；'.join(M['caveats'])))
    r = 3
    for k, v in notes:
        if k == '' and v == '':
            r += 1
            continue
        if v == '':
            c = ws8.cell(r, 1, k)
            c.font = SUB
            ws8.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            for j in (1, 2):
                ws8.cell(r, j).fill = PatternFill('solid', fgColor='DDEBF7')
                ws8.cell(r, j).border = BD
        else:
            ws8.cell(r, 1, k).font = CELL
            ws8.cell(r, 1).alignment = LT
            ws8.cell(r, 2, v).font = NOTE
            ws8.cell(r, 2).alignment = LT
            for j in (1, 2):
                ws8.cell(r, j).border = BD
        r += 1
    ws8.column_dimensions['A'].width = 24
    ws8.column_dimensions['B'].width = 110

    # ⑨ 数据溯源：每个聚合指标都可回算到明细，原声还保留来源行。
    trace = S.get('trace') or {}
    ws9 = wb.create_sheet('⑨ 数据溯源')
    ws9['A1'] = '数据血缘与复核信息'
    ws9['A1'].font = TITLE
    trace_rows = [
        ('技能版本', M.get('skill_version', 'unknown')),
        ('源数据文件', M.get('data_file', '')),
        ('源数据 SHA-256', trace.get('source_file_sha256', '')),
        ('源数据行数', trace.get('source_rows', '')),
        ('主口径行数', trace.get('main_scope_rows', '')),
        ('口径配置 SHA-256', trace.get('scope_sha256', '')),
        ('明细定位规则', trace.get('row_locator', '')),
        ('复核方法', '以②主口径明细、⑤外溢明细的“来源定位”回指 CSV；统计指标由这些明细按⑧口径与方法说明复算。'),
        ('证据边界', '本报告只陈述用户提供数据中的统计事实；待验证方向不构成因果结论。'),
    ]
    table(ws9, 3, ['项目', '内容'], trace_rows, widths=[26, 118], wrap=(2,))

    wb.save(a.out)
    print(f'OK -> {a.out}')
    print('sheets:', wb.sheetnames)

if __name__ == '__main__':
    main()
